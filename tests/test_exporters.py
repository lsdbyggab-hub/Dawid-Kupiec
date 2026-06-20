import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from wall_pdf_analyzer.analyzer import analyze_project
from wall_pdf_analyzer.exporters import export_control_pdf, export_control_svg, export_result
from wall_pdf_analyzer.io import load_project
from wall_pdf_analyzer.models import AnalysisInput, AnalysisScope, Point, Scale, WallSegment, WallType


class ExporterTests(unittest.TestCase):
    def test_exports_json_xlsx_and_svg(self):
        project = load_project("examples/sample_project.json")
        result = analyze_project(project)

        with tempfile.TemporaryDirectory() as temp_dir:
            base = Path(temp_dir)
            json_path = base / "report.json"
            xlsx_path = base / "report.xlsx"
            svg_path = base / "overlay.svg"
            pdf_path = base / "overlay.pdf"

            export_result(result, json_path)
            export_result(result, xlsx_path)
            export_control_svg(project, result, svg_path)
            export_control_pdf(project, result, pdf_path)

            payload = json.loads(json_path.read_text(encoding="utf-8"))
            workbook = load_workbook(xlsx_path)
            svg = svg_path.read_text(encoding="utf-8")
            pdf_size = pdf_path.stat().st_size

        self.assertEqual(payload["project_name"], result.project_name)
        self.assertIn("Podsumowanie", workbook.sheetnames)
        self.assertIn("Odcinki", workbook.sheetnames)
        self.assertIn("#FF00FF", svg)
        self.assertIn('stroke-width="1.2"', svg)
        self.assertNotIn('stroke-width="7"', svg)
        self.assertNotIn('stroke-width="5"', svg)
        self.assertGreater(pdf_size, 0)

    def test_control_pdf_splits_large_plan_into_detail_pages(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            base = Path(temp_dir)
            source_pdf = base / "large-plan.pdf"
            output_pdf = base / "large-overlay.pdf"
            _write_blank_pdf(source_pdf, width=1400, height=1000)
            project = _large_project(source_pdf)
            result = analyze_project(project)

            export_control_pdf(project, result, output_pdf)

            import fitz  # type: ignore

            document = fitz.open(str(output_pdf))
            try:
                page_count = len(document)
            finally:
                document.close()

        self.assertGreater(page_count, 1)

def _write_blank_pdf(path: Path, *, width: float, height: float) -> None:
    import fitz  # type: ignore

    document = fitz.open()
    page = document.new_page(width=width, height=height)
    page.insert_text((40, 40), "Large plan", fontsize=12)
    document.save(str(path))
    document.close()


def _large_project(source_pdf: Path) -> AnalysisInput:
    wall_type = WallType(code="IV20", name="Sciana IV20", color="#4C1D95")
    segments = tuple(
        WallSegment(
            id=f"W-{index:03d}",
            type_code="IV20",
            start=Point(80 + index * 25, 120),
            end=Point(80 + index * 25, 880),
            page=1,
            confidence=0.9,
            measurement_basis="test",
        )
        for index in range(40)
    )
    return AnalysisInput(
        project_name="large-plan",
        scale=Scale(drawing_units_per_meter=28.346, label="1:100"),
        scope=AnalysisScope(pages=(1,)),
        wall_types={"IV20": wall_type},
        wall_segments=segments,
        source_metadata={
            "source_pdf": str(source_pdf),
            "coordinate_origin": "pdf_bottom_left",
        },
    )


if __name__ == "__main__":
    unittest.main()
