from __future__ import annotations

import csv
import json
from dataclasses import asdict
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from wall_pdf_analyzer.models import AnalysisInput, AnalysisResult

SVG_WALL_LINE_WIDTH = 0.9
SVG_EXTERIOR_LINE_WIDTH = 1.2
PDF_WALL_LINE_WIDTH = 0.65
PDF_EXTERIOR_LINE_WIDTH = 0.85
PDF_OPENING_MARK_RADIUS = 3.0
PDF_TILE_SOURCE_SIDE = 650.0
PDF_TILE_MAX_GRID = 4
PDF_TILE_OVERLAP = 18.0
PDF_TILE_PAGE_WIDTH = 842.0
PDF_TILE_PAGE_HEIGHT = 595.0
PDF_TILE_MARGIN = 24.0
PDF_TILE_HEADER_HEIGHT = 42.0


def export_result(result: AnalysisResult, path: str | Path) -> None:
    destination = Path(path)
    suffix = destination.suffix.lower()
    if suffix == ".xlsx":
        export_xlsx(result, destination)
    elif suffix == ".csv":
        export_csv(result, destination)
    elif suffix == ".json":
        export_json(result, destination)
    else:
        raise ValueError("Supported export formats: .xlsx, .csv, .json")


def export_json(result: AnalysisResult, path: str | Path) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "project_name": result.project_name,
        "scale": result.scale_label,
        "scope": result.scope,
        "source_metadata": result.source_metadata,
        "measurement_method": result.source_metadata.get("method", ""),
        "exterior_total_m": result.exterior_total_m,
        "summary": [asdict(row) for row in result.summary],
        "rows": [asdict(row) for row in result.rows],
        "wall_measurements": [_wall_measurement_payload(row) for row in result.rows],
        "assumptions": list(result.source_metadata.get("assumptions", [])),
        "warnings": list(result.warnings),
        "professional_note": (
            "Pomiary z rysunkow nalezy zweryfikowac z oryginalnymi plikami "
            "CAD/BIM, oficjalnymi wymiarami albo pomiarem na miejscu przed "
            "zamawianiem materialow, wycena lub wykonaniem prac."
        ),
    }
    if "swedish_plan_context" in result.source_metadata:
        payload["swedish_plan_context"] = result.source_metadata["swedish_plan_context"]
    if "wall_type_legend" in result.source_metadata:
        payload["wall_type_legend"] = result.source_metadata["wall_type_legend"]
    destination.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def export_csv(result: AnalysisResult, path: str | Path) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "segment_id",
        "wall_type",
        "wall_name",
        "color",
        "page",
        "gross_length_m",
        "openings_m",
        "opening_count",
        "opening_kinds",
        "net_length_m",
        "confidence",
        "exterior",
        "scope",
        "measurement_basis",
        "centerline_or_face",
        "comment",
    ]
    with destination.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        for row in result.rows:
            writer.writerow(asdict(row))


def export_xlsx(result: AnalysisResult, path: str | Path) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Podsumowanie"
    summary_sheet.append(["Projekt", result.project_name])
    summary_sheet.append(["Skala", result.scale_label])
    summary_sheet.append(["Zakres", result.scope])
    summary_sheet.append(["Sciany zewnetrzne brutto [m]", result.exterior_total_m])
    if result.source_metadata:
        summary_sheet.append(["Typ PDF", result.source_metadata.get("pdf_kind", "")])
        summary_sheet.append(["Metoda", result.source_metadata.get("method", "")])
        summary_sheet.append(["Zrodlo skali", result.source_metadata.get("scale_source", "")])
    summary_sheet.append([])
    header_row = summary_sheet.max_row + 1
    summary_sheet.append(
        [
            "Typ",
            "Nazwa",
            "Kolor",
            "Liczba odcinkow",
            "Brutto [m]",
            "Otwory [m]",
            "Liczba otworow",
            "Typy otworow",
            "Netto [m]",
            "Zewnetrzna",
        ]
    )
    _style_header(summary_sheet, header_row)
    for row in result.summary:
        summary_sheet.append(
            [
                row.wall_type,
                row.wall_name,
                row.color,
                row.segment_count,
                row.gross_length_m,
                row.openings_m,
                row.opening_count,
                row.opening_kinds,
                row.net_length_m,
                "tak" if row.exterior else "nie",
            ]
        )
        _paint_color_cell(summary_sheet.cell(summary_sheet.max_row, 3), row.color)

    detail_sheet = workbook.create_sheet("Odcinki")
    detail_sheet.append(
        [
            "ID",
            "Typ",
            "Nazwa",
            "Kolor",
            "Strona",
            "Brutto [m]",
            "Otwory [m]",
            "Liczba otworow",
            "Typy otworow",
            "Netto [m]",
            "Pewnosc",
            "Zewnetrzna",
            "Zakres",
            "Podstawa pomiaru",
            "Os/krawedz",
            "Komentarz",
        ]
    )
    _style_header(detail_sheet, 1)
    for row in result.rows:
        detail_sheet.append(
            [
                row.segment_id,
                row.wall_type,
                row.wall_name,
                row.color,
                row.page,
                row.gross_length_m,
                row.openings_m,
                row.opening_count,
                row.opening_kinds,
                row.net_length_m,
                row.confidence,
                "tak" if row.exterior else "nie",
                row.scope,
                row.measurement_basis,
                row.centerline_or_face,
                row.comment,
            ]
        )
        _paint_color_cell(detail_sheet.cell(detail_sheet.max_row, 4), row.color)

    audit_sheet = workbook.create_sheet("Kontrola")
    audit_sheet.append(["Ostrzezenia"])
    _style_header(audit_sheet, 1)
    if result.warnings:
        for warning in result.warnings:
            audit_sheet.append([warning])
    else:
        audit_sheet.append(["Brak ostrzezen"])
    assumptions = result.source_metadata.get("assumptions", [])
    if assumptions:
        audit_sheet.append([])
        audit_sheet.append(["Zalozenia"])
        _style_header(audit_sheet, audit_sheet.max_row)
        for assumption in assumptions:
            audit_sheet.append([assumption])

    for sheet in workbook.worksheets:
        _autosize(sheet)

    workbook.save(destination)


def _wall_measurement_payload(row) -> dict[str, object]:
    return {
        "segment_id": row.segment_id,
        "page": row.page,
        "wall_type": row.wall_type,
        "length": row.gross_length_m,
        "unit": "m",
        "thickness": None,
        "adjacent_room": "",
        "openings_width_m": row.openings_m,
        "opening_count": row.opening_count,
        "opening_kinds": row.opening_kinds,
        "net_length_m": row.net_length_m,
        "openings_deducted": row.openings_m > 0,
        "confidence": row.confidence,
        "measurement_basis": row.measurement_basis,
        "centerline_or_face": row.centerline_or_face,
        "notes": row.comment,
    }


def export_control_svg(
    project: AnalysisInput,
    result: AnalysisResult,
    path: str | Path,
) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    rows_by_id = {row.segment_id: row for row in result.rows}
    segments = [
        segment for segment in project.wall_segments if segment.id in rows_by_id
    ]
    min_x = min(min(segment.start.x, segment.end.x) for segment in segments)
    min_y = min(min(segment.start.y, segment.end.y) for segment in segments)
    max_x = max(max(segment.start.x, segment.end.x) for segment in segments)
    max_y = max(max(segment.start.y, segment.end.y) for segment in segments)
    padding = 30
    width = max(max_x - min_x + padding * 2, 120)
    height = max(max_y - min_y + padding * 2, 120)

    def tx(value: float) -> float:
        return value - min_x + padding

    def ty(value: float) -> float:
        return value - min_y + padding

    lines = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        (
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{width:.0f}" '
            f'height="{height:.0f}" viewBox="0 0 {width:.0f} {height:.0f}">'
        ),
        '<rect width="100%" height="100%" fill="#F8FAFC"/>',
        (
            f'<text x="16" y="24" font-family="Arial" font-size="14" '
            f'fill="#111827">{escape(result.project_name)}</text>'
        ),
    ]
    for segment in segments:
        row = rows_by_id[segment.id]
        x1, y1 = tx(segment.start.x), ty(segment.start.y)
        x2, y2 = tx(segment.end.x), ty(segment.end.y)
        stroke_width = SVG_EXTERIOR_LINE_WIDTH if row.exterior else SVG_WALL_LINE_WIDTH
        lines.append(
            f'<line x1="{x1:.2f}" y1="{y1:.2f}" x2="{x2:.2f}" y2="{y2:.2f}" '
            f'stroke="{escape(row.color)}" stroke-width="{stroke_width}" '
            'stroke-linecap="round"/>'
        )
        label_x = (x1 + x2) / 2
        label_y = (y1 + y2) / 2 - 8
        lines.append(
            f'<text x="{label_x:.2f}" y="{label_y:.2f}" font-family="Arial" '
            f'font-size="12" fill="#111827">{escape(row.wall_type)} '
            f'{row.gross_length_m:.2f} m</text>'
        )
        for opening in segment.openings:
            ratio = 0.5 if opening.offset is None else max(
                0.0, min(opening.offset / max(segment.drawing_length, 1), 1.0)
            )
            ox = x1 + (x2 - x1) * ratio
            oy = y1 + (y2 - y1) * ratio
            lines.append(
                f'<circle cx="{ox:.2f}" cy="{oy:.2f}" r="5" fill="#FFFFFF" '
                f'stroke="#111827" stroke-width="1.5"/>'
            )
            lines.append(
                f'<text x="{ox + 7:.2f}" y="{oy - 7:.2f}" font-family="Arial" '
                f'font-size="10" fill="#111827">{escape(opening.kind)}</text>'
            )
    lines.append("</svg>")
    destination.write_text("\n".join(lines), encoding="utf-8")


def export_control_pdf(
    project: AnalysisInput,
    result: AnalysisResult,
    path: str | Path,
) -> None:
    try:
        import fitz  # type: ignore
    except ImportError as exc:  # pragma: no cover - depends on user environment
        raise RuntimeError("Do eksportu kontrolnego PDF potrzebna jest biblioteka PyMuPDF.") from exc
    try:
        fitz.TOOLS.mupdf_display_errors(False)
        fitz.TOOLS.mupdf_display_warnings(False)
    except Exception:
        pass

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    source_pdf = result.source_metadata.get("source_pdf")
    if source_pdf and Path(str(source_pdf)).exists():
        try:
            document = fitz.open(str(source_pdf))
            try:
                _draw_pdf_overlay(document, project, result)
                output_document = fitz.open()
                try:
                    _copy_pdf_with_detail_tiles(output_document, document, project, result)
                    output_document.save(str(destination))
                finally:
                    output_document.close()
                return
            finally:
                document.close()
        except Exception:
            pass

    document = fitz.open()
    try:
        blank_overlay = fitz.open()
        try:
            _draw_blank_control_pdf(blank_overlay, project, result)
            _copy_pdf_with_detail_tiles(document, blank_overlay, project, result)
        finally:
            blank_overlay.close()
        document.save(str(destination))
    finally:
        document.close()


def _copy_pdf_with_detail_tiles(
    target_document,
    overlay_document,
    project: AnalysisInput,
    result: AnalysisResult,
) -> None:
    segments_by_page = _segments_by_page(project, result)
    for page_index in range(len(overlay_document)):
        target_document.insert_pdf(overlay_document, from_page=page_index, to_page=page_index)
        source_page = overlay_document[page_index]
        page_number = page_index + 1
        columns, rows = _tile_grid(source_page.rect, len(segments_by_page.get(page_number, [])))
        if columns == 1 and rows == 1:
            continue
        _append_detail_tile_pages(
            target_document,
            overlay_document,
            page_index,
            columns=columns,
            rows=rows,
        )


def _segments_by_page(project: AnalysisInput, result: AnalysisResult) -> dict[int, list]:
    rows_by_id = {row.segment_id for row in result.rows}
    pages: dict[int, list] = {}
    for segment in project.wall_segments:
        if segment.id in rows_by_id:
            pages.setdefault(segment.page, []).append(segment)
    if not pages and project.wall_segments:
        pages[1] = list(project.wall_segments)
    return pages


def _tile_grid(page_rect, segment_count: int) -> tuple[int, int]:
    if segment_count <= 0:
        return 1, 1
    columns = max(1, min(PDF_TILE_MAX_GRID, int((page_rect.width + PDF_TILE_SOURCE_SIDE - 1) // PDF_TILE_SOURCE_SIDE)))
    rows = max(1, min(PDF_TILE_MAX_GRID, int((page_rect.height + PDF_TILE_SOURCE_SIDE - 1) // PDF_TILE_SOURCE_SIDE)))
    if segment_count >= 35 and columns == 1 and rows == 1:
        if page_rect.width >= page_rect.height:
            columns = 2
        else:
            rows = 2
    if segment_count >= 80:
        columns = max(columns, 2)
        rows = max(rows, 2)
    return columns, rows


def _append_detail_tile_pages(
    target_document,
    overlay_document,
    page_index: int,
    *,
    columns: int,
    rows: int,
) -> None:
    import fitz  # type: ignore

    source_page = overlay_document[page_index]
    source_rect = source_page.rect
    tile_width = source_rect.width / columns
    tile_height = source_rect.height / rows
    target_rect = fitz.Rect(
        PDF_TILE_MARGIN,
        PDF_TILE_HEADER_HEIGHT,
        PDF_TILE_PAGE_WIDTH - PDF_TILE_MARGIN,
        PDF_TILE_PAGE_HEIGHT - PDF_TILE_MARGIN,
    )
    total = columns * rows
    tile_number = 1
    for row_index in range(rows):
        for column_index in range(columns):
            clip = fitz.Rect(
                source_rect.x0 + column_index * tile_width,
                source_rect.y0 + row_index * tile_height,
                source_rect.x0 + (column_index + 1) * tile_width,
                source_rect.y0 + (row_index + 1) * tile_height,
            )
            clip = _expanded_clip(clip, source_rect, PDF_TILE_OVERLAP)
            page = target_document.new_page(width=PDF_TILE_PAGE_WIDTH, height=PDF_TILE_PAGE_HEIGHT)
            page.insert_text(
                (PDF_TILE_MARGIN, 24),
                f"Strona {page_index + 1} - fragment {tile_number}/{total}",
                fontsize=10,
                color=(0.07, 0.09, 0.15),
            )
            page.show_pdf_page(target_rect, overlay_document, page_index, clip=clip)
            page.draw_rect(target_rect, color=(0.65, 0.70, 0.78), width=0.35)
            tile_number += 1


def _expanded_clip(clip, page_rect, overlap: float):
    import fitz  # type: ignore

    return fitz.Rect(
        max(page_rect.x0, clip.x0 - overlap),
        max(page_rect.y0, clip.y0 - overlap),
        min(page_rect.x1, clip.x1 + overlap),
        min(page_rect.y1, clip.y1 + overlap),
    )


def _draw_pdf_overlay(document, project: AnalysisInput, result: AnalysisResult) -> None:
    rows_by_id = {row.segment_id: row for row in result.rows}
    origin = result.source_metadata.get("coordinate_origin", "image_top_left")
    for segment in project.wall_segments:
        row = rows_by_id.get(segment.id)
        if row is None or segment.page < 1 or segment.page > len(document):
            continue
        page = document[segment.page - 1]
        p1 = _pdf_point(page, segment.start.x, segment.start.y, origin)
        p2 = _pdf_point(page, segment.end.x, segment.end.y, origin)
        color = _pdf_color(row.color)
        page.draw_line(
            p1,
            p2,
            color=color,
            width=PDF_EXTERIOR_LINE_WIDTH if row.exterior else PDF_WALL_LINE_WIDTH,
            overlay=True,
        )
        label = f"{row.wall_type} {row.gross_length_m:.2f} m"
        midpoint = ((p1.x + p2.x) / 2, (p1.y + p2.y) / 2)
        page.insert_text(
            midpoint,
            label,
            fontsize=6,
            color=color,
            overlay=True,
        )
        _draw_pdf_openings(page, segment, row, origin)


def _draw_blank_control_pdf(document, project: AnalysisInput, result: AnalysisResult) -> None:
    rows_by_id = {row.segment_id: row for row in result.rows}
    segments = [segment for segment in project.wall_segments if segment.id in rows_by_id]
    if not segments:
        document.new_page(width=595, height=842)
        return
    min_x = min(min(segment.start.x, segment.end.x) for segment in segments)
    min_y = min(min(segment.start.y, segment.end.y) for segment in segments)
    max_x = max(max(segment.start.x, segment.end.x) for segment in segments)
    max_y = max(max(segment.start.y, segment.end.y) for segment in segments)
    padding = 36
    width = max(max_x - min_x + padding * 2, 300)
    height = max(max_y - min_y + padding * 2, 220)
    page = document.new_page(width=width, height=height)
    page.insert_text((16, 24), result.project_name, fontsize=12, color=(0.07, 0.09, 0.15))

    def tx(value: float) -> float:
        return value - min_x + padding

    def ty(value: float) -> float:
        return value - min_y + padding

    for segment in segments:
        row = rows_by_id[segment.id]
        p1 = (tx(segment.start.x), ty(segment.start.y))
        p2 = (tx(segment.end.x), ty(segment.end.y))
        color = _pdf_color(row.color)
        page.draw_line(
            p1,
            p2,
            color=color,
            width=PDF_EXTERIOR_LINE_WIDTH if row.exterior else PDF_WALL_LINE_WIDTH,
            overlay=True,
        )
        page.insert_text(
            ((p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2 - 4),
            f"{row.wall_type} {row.gross_length_m:.2f} m",
            fontsize=6,
            color=color,
            overlay=True,
        )
        _draw_blank_pdf_openings(page, segment, row, tx, ty)


def _pdf_point(page, x: float, y: float, origin: str):
    import fitz  # type: ignore

    if origin == "pdf_bottom_left":
        return fitz.Point(x, page.rect.height - y)
    return fitz.Point(x, y)


def _draw_pdf_openings(page, segment, row, origin: str) -> None:
    for opening in segment.openings:
        point = _opening_pdf_point(page, segment, opening, origin)
        page.draw_circle(
            point,
            PDF_OPENING_MARK_RADIUS,
            color=(0.07, 0.09, 0.15),
            fill=(1, 1, 1),
            width=0.5,
            overlay=True,
        )
        page.insert_text(
            (point.x + 5, point.y - 5),
            opening.kind,
            fontsize=6,
            color=(0.07, 0.09, 0.15),
            overlay=True,
        )


def _draw_blank_pdf_openings(page, segment, row, tx, ty) -> None:
    import fitz  # type: ignore

    for opening in segment.openings:
        ratio = 0.5 if opening.offset is None else max(
            0.0,
            min(opening.offset / max(segment.drawing_length, 1), 1.0),
        )
        x = tx(segment.start.x + (segment.end.x - segment.start.x) * ratio)
        y = ty(segment.start.y + (segment.end.y - segment.start.y) * ratio)
        point = fitz.Point(x, y)
        page.draw_circle(
            point,
            PDF_OPENING_MARK_RADIUS,
            color=(0.07, 0.09, 0.15),
            fill=(1, 1, 1),
            width=0.5,
            overlay=True,
        )
        page.insert_text((x + 5, y - 5), opening.kind, fontsize=6, color=(0.07, 0.09, 0.15), overlay=True)


def _opening_pdf_point(page, segment, opening, origin: str):
    ratio = 0.5 if opening.offset is None else max(
        0.0,
        min(opening.offset / max(segment.drawing_length, 1), 1.0),
    )
    x = segment.start.x + (segment.end.x - segment.start.x) * ratio
    y = segment.start.y + (segment.end.y - segment.start.y) * ratio
    return _pdf_point(page, x, y, origin)


def _pdf_color(color: str) -> tuple[float, float, float]:
    if color.startswith("#") and len(color) == 7:
        return (
            int(color[1:3], 16) / 255,
            int(color[3:5], 16) / 255,
            int(color[5:7], 16) / 255,
        )
    return (0.1, 0.2, 0.8)


def _style_header(sheet, row_number: int) -> None:
    for cell in sheet[row_number]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E5E7EB")


def _paint_color_cell(cell, color: str) -> None:
    cell.value = color
    if color.startswith("#") and len(color) == 7:
        cell.fill = PatternFill("solid", fgColor=color[1:].upper())


def _autosize(sheet) -> None:
    for column_cells in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        sheet.column_dimensions[column_letter].width = min(max_length + 2, 60)
